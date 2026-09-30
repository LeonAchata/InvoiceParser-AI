"""Word (.docx), Excel (.xlsx) and CSV/TSV documents -> plain text (tables kept as rows)."""

import csv
from io import BytesIO, StringIO

from PIL import Image

from app.loaders.base import LoadedDocument, UnsupportedFormatError
from app.loaders.image import normalize_image
from app.loaders.text import decode_text

# If a DOCX has less text than this, embedded pictures (e.g. a pasted photo of a receipt) are sent too.
DOCX_MIN_TEXT = 200


def load_docx(data: bytes, fmt: str, *, max_pages: int, max_side: int, **_) -> LoadedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise UnsupportedFormatError(f"Invalid DOCX file: {exc}") from exc

    lines: list[str] = []

    def table_rows(table: Table) -> None:
        for row in table.rows:
            cells: list[str] = []
            seen = []  # keep element references (not ids) so identity checks stay valid
            for cell in row.cells:
                if any(cell._tc is tc for tc in seen):  # merged cells are repeated once per grid column
                    continue
                seen.append(cell._tc)
                cells.append(cell.text.strip())
            if any(cells):
                lines.append(" | ".join(cells))

    # Walk the body in order so paragraphs and tables stay interleaved as in the document.
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            table_rows(Table(child, doc))

    for section in doc.sections:  # company data often lives in headers/footers
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                if p.text.strip():
                    lines.append(p.text.strip())
            for t in part.tables:
                table_rows(t)

    text = "\n".join(lines)
    images = []
    if len(text) < DOCX_MIN_TEXT:
        for rel in doc.part.rels.values():
            if "image" in rel.reltype and len(images) < max_pages:
                try:
                    images.append(normalize_image(Image.open(BytesIO(rel.target_part.blob)), max_side))
                except Exception:
                    continue

    method = "native-text" + (" + vision" if images else "")
    return LoadedDocument(format=fmt, text=text, images=images, method=method, meta={"paragraphs": len(lines)})


def load_xlsx(data: bytes, fmt: str, **_) -> LoadedDocument:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise UnsupportedFormatError(f"Invalid Excel file: {exc}") from exc

    blocks = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"--- Sheet: {ws.title} ---\n" + "\n".join(rows))
    wb.close()

    return LoadedDocument(format=fmt, text="\n\n".join(blocks), pages=len(blocks) or 1, method="native-text")


def load_csv(data: bytes, fmt: str, **_) -> LoadedDocument:
    content = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if fmt == "tsv" else csv.excel
    rows = [" | ".join(c.strip() for c in row) for row in csv.reader(StringIO(content), dialect) if any(row)]
    return LoadedDocument(format=fmt, text="\n".join(rows), method="native-text")
