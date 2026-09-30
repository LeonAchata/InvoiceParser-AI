"""Generate fictitious sample documents in every supported family of formats.

    python scripts/generate_samples.py

All companies, people and tax ids are made up (RUCs have a valid check digit so validation passes).
"""

import random
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageFilter

OUT = Path(__file__).resolve().parent.parent / "samples"


def ruc(base: str) -> str:
    weights = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
    digit = 11 - sum(int(d) * w for d, w in zip(base, weights)) % 11
    return base + str({10: 0, 11: 1}.get(digit, digit))


ISSUER_RUC = ruc("2060987654")
CUSTOMER_RUC = ruc("2055512345")

FACTURA_HTML = f"""
<div style="font-family: sans-serif; font-size: 10px; color: #222;">
  <table style="width:100%"><tr>
    <td>
      <p style="font-size:18px; font-weight:bold; color:#1e3a8a">ANDES CLOUD SOLUTIONS S.A.C.</p>
      <p>Av. José Larco 1232, Of. 501 - Miraflores - Lima<br>Tel. (01) 555-0199 · facturacion@andescloud.pe</p>
    </td>
    <td style="border:2px solid #1e3a8a; text-align:center; padding:8px; width:170px">
      <b>R.U.C. {ISSUER_RUC}</b><br><b style="font-size:13px">FACTURA ELECTRÓNICA</b><br><b>F001-00004821</b>
    </td>
  </tr></table>
  <br>
  <table style="width:100%; border:1px solid #999; padding:4px">
    <tr><td><b>Fecha de emisión:</b> 14/08/2026</td><td><b>Fecha de vencimiento:</b> 13/09/2026</td></tr>
    <tr><td><b>Señor(es):</b> COMERCIAL LOS PORTALES DEL SUR E.I.R.L.</td><td><b>RUC:</b> {CUSTOMER_RUC}</td></tr>
    <tr><td colspan="2"><b>Dirección:</b> Calle Los Cedros 245 Urb. El Remanso - La Molina - Lima</td></tr>
    <tr><td><b>Forma de pago:</b> Crédito</td><td><b>Moneda:</b> SOLES</td></tr>
  </table>
  <br>
  <table style="width:100%; border-collapse:collapse" border="1" cellpadding="4">
    <tr style="background:#1e3a8a; color:white"><th>Cant.</th><th>Unidad</th><th>Descripción</th><th>P. Unit.</th><th>Importe</th></tr>
    <tr><td>1</td><td>SERV</td><td>Implementación de plataforma ERP en la nube (fase 1)</td><td>4,500.00</td><td>4,500.00</td></tr>
    <tr><td>12</td><td>MES</td><td>Licencia mensual usuario ERP</td><td>85.00</td><td>1,020.00</td></tr>
    <tr><td>8</td><td>HORA</td><td>Capacitación presencial a usuarios</td><td>60.00</td><td>480.00</td></tr>
  </table>
  <br>
  <table style="width:100%"><tr><td style="width:55%">
      SON: SIETE MIL OCHENTA Y 00/100 SOLES<br><br>
      <b>Operación sujeta al Sistema de Pago de Obligaciones Tributarias (SPOT)</b><br>
      Detracción 12%: S/ 850.00 · Cta. Banco de la Nación 00-000-123456
    </td><td>
      <table style="width:100%">
        <tr><td>Op. Gravada</td><td style="text-align:right">S/ 6,000.00</td></tr>
        <tr><td>IGV 18%</td><td style="text-align:right">S/ 1,080.00</td></tr>
        <tr><td><b>Importe Total</b></td><td style="text-align:right"><b>S/ 7,080.00</b></td></tr>
      </table>
  </td></tr></table>
  <p style="color:#666; font-size:8px">Representación impresa de la Factura Electrónica. Documento ficticio de ejemplo.</p>
</div>
"""

BOLETA_HTML = f"""
<div style="font-family: monospace; font-size: 11px; text-align:center">
  <b style="font-size:14px">BODEGA DOÑA ROSA</b><br>
  de Rosa Quispe Mamani<br>RUC {ruc("1045678901")}<br>Jr. Ayacucho 318 - Cercado - Arequipa<br>
  --------------------------------<br>
  <b>BOLETA DE VENTA ELECTRÓNICA</b><br>B002-00018733<br>
  Fecha: 03/09/2026  Hora: 19:42<br>
  Cliente: JUAN CARLOS TORRES VEGA<br>DNI: 45781236<br>
  --------------------------------<br>
  <table style="width:100%; font-family: monospace; font-size: 11px">
    <tr><td style="text-align:left">2 x Leche Gloria 1L</td><td style="text-align:right">11.80</td></tr>
    <tr><td style="text-align:left">1 x Arroz Costeño 5kg</td><td style="text-align:right">24.90</td></tr>
    <tr><td style="text-align:left">3 x Pan francés (bolsa)</td><td style="text-align:right">6.00</td></tr>
    <tr><td style="text-align:left">1 x Aceite Primor 900ml</td><td style="text-align:right">10.50</td></tr>
  </table>
  --------------------------------<br>
  <table style="width:100%; font-family: monospace; font-size: 11px">
    <tr><td style="text-align:left">OP. GRAVADA</td><td style="text-align:right">45.08</td></tr>
    <tr><td style="text-align:left">IGV 18%</td><td style="text-align:right">8.12</td></tr>
    <tr><td style="text-align:left"><b>TOTAL S/</b></td><td style="text-align:right"><b>53.20</b></td></tr>
  </table>
  Pago: YAPE<br>¡Gracias por su compra!
</div>
"""


def html_to_pdf(html: str, size: tuple[float, float]) -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=size[0], height=size[1])
    page.insert_htmlbox(fitz.Rect(36, 36, size[0] - 36, size[1] - 36), html)
    return doc.tobytes(deflate=True)


def scanned_look(pdf_bytes: bytes, dpi: int = 160, angle: float = 1.4) -> Image.Image:
    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        pix = doc[0].get_pixmap(dpi=dpi)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L")
    rng = random.Random(7)
    noise = Image.effect_noise(img.size, 18).point(lambda p: 255 if p > 110 else 235)
    img = Image.composite(img, noise, img.point(lambda p: 255 if p < 200 else 0))
    img = img.rotate(angle, expand=True, fillcolor=rng.randint(225, 240)).filter(ImageFilter.GaussianBlur(0.6))
    return img.convert("RGB")


def make_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    title = doc.add_paragraph()
    run = title.add_run("NORTHWIND DESIGN STUDIO LLC")
    run.bold, run.font.size, run.font.color.rgb = True, Pt(18), RGBColor(0x0F, 0x76, 0x6E)
    doc.add_paragraph("742 Evergreen Ave, Suite 12, Austin, TX 78701 · EIN 84-2918375")
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    heading.add_run("INVOICE #INV-2026-0192").bold = True
    doc.add_paragraph(
        "Invoice date: September 2, 2026\nDue date: October 2, 2026\nPayment terms: Bank transfer, Net 30"
    )
    doc.add_paragraph("Bill to:\nPacific Harbor Foods Inc.\nTax ID 91-4410226\n1800 Harbor Blvd, San Diego, CA 92101")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Description", "Qty", "Rate", "Amount")):
        cell.text = text
    for row in (
        ("Brand identity refresh (logo, palette, typography)", "1", "$3,200.00", "$3,200.00"),
        ("Packaging design - SKU variants", "6", "$450.00", "$2,700.00"),
        ("Photography retouching (per image)", "40", "$12.50", "$500.00"),
    ):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    doc.add_paragraph()
    totals = doc.add_paragraph(
        "Subtotal: $6,400.00\nDiscount (5%): -$320.00\nSales tax (8.25%): $501.60\nTOTAL DUE: $6,581.60"
    )
    totals.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("Thank you for your business! Fictitious sample document.")
    doc.save(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Factura"
    rows = [
        ["INVERSIONES MAR AZUL S.A.C.", None, None, "RUC", ruc("2071122334")],
        ["Av. Grau 1020 - Chorrillos - Lima"],
        [],
        ["FACTURA", "F003-00000077", None, "Fecha", "2026-07-21"],
        ["Cliente", "HOTEL COSTA VERDE S.A.", None, "RUC", ruc("2051234567")],
        ["Dirección", "Malecón Pardo 450 - Miraflores"],
        ["Moneda", "USD", None, "Pago", "Transferencia"],
        [],
        ["Descripción", "Cantidad", "P. Unitario", "Importe"],
        ["Mantenimiento preventivo de aire acondicionado", 10, 45, 450],
        ["Reemplazo de filtros HEPA", 20, 18.5, 370],
        [],
        [None, None, "Subtotal", 820],
        [None, None, "IGV 18%", 147.6],
        [None, None, "Total", 967.6],
    ]
    for row in rows:
        ws.append(row)
    for cell in ("A1", "A4", "B4", "C15", "D15"):
        ws[cell].font = Font(bold=True)
    ws.column_dimensions["A"].width = 46
    wb.save(path)


def make_xml(path: Path) -> None:
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
         xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
         xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2"
         xmlns:ds="http://www.w3.org/2000/09/xmldsig#">
  <ds:Signature Id="SignSUNAT"><ds:SignatureValue>QmFzZTY0U2lnbmF0dXJlUGxhY2Vob2xkZXI=</ds:SignatureValue></ds:Signature>
  <cbc:UBLVersionID>2.1</cbc:UBLVersionID>
  <cbc:ID>F010-00000315</cbc:ID>
  <cbc:IssueDate>2026-09-10</cbc:IssueDate>
  <cbc:DueDate>2026-10-10</cbc:DueDate>
  <cbc:InvoiceTypeCode listID="0101">01</cbc:InvoiceTypeCode>
  <cbc:DocumentCurrencyCode>PEN</cbc:DocumentCurrencyCode>
  <cac:AccountingSupplierParty><cac:Party>
    <cac:PartyIdentification><cbc:ID schemeID="6">{ISSUER_RUC}</cbc:ID></cac:PartyIdentification>
    <cac:PartyLegalEntity><cbc:RegistrationName>ANDES CLOUD SOLUTIONS S.A.C.</cbc:RegistrationName>
      <cac:RegistrationAddress><cbc:District>MIRAFLORES</cbc:District>
        <cac:AddressLine><cbc:Line>AV. JOSE LARCO 1232 OF. 501</cbc:Line></cac:AddressLine></cac:RegistrationAddress>
    </cac:PartyLegalEntity></cac:Party></cac:AccountingSupplierParty>
  <cac:AccountingCustomerParty><cac:Party>
    <cac:PartyIdentification><cbc:ID schemeID="6">{CUSTOMER_RUC}</cbc:ID></cac:PartyIdentification>
    <cac:PartyLegalEntity><cbc:RegistrationName>COMERCIAL LOS PORTALES DEL SUR E.I.R.L.</cbc:RegistrationName>
      <cac:RegistrationAddress><cbc:District>LA MOLINA</cbc:District>
        <cac:AddressLine><cbc:Line>CALLE LOS CEDROS 245 URB. EL REMANSO</cbc:Line></cac:AddressLine></cac:RegistrationAddress>
    </cac:PartyLegalEntity></cac:Party></cac:AccountingCustomerParty>
  <cac:PaymentTerms><cbc:ID>FormaPago</cbc:ID><cbc:PaymentMeansID>Contado</cbc:PaymentMeansID></cac:PaymentTerms>
  <cac:TaxTotal><cbc:TaxAmount currencyID="PEN">324.00</cbc:TaxAmount></cac:TaxTotal>
  <cac:LegalMonetaryTotal>
    <cbc:LineExtensionAmount currencyID="PEN">1800.00</cbc:LineExtensionAmount>
    <cbc:PayableAmount currencyID="PEN">2124.00</cbc:PayableAmount>
  </cac:LegalMonetaryTotal>
  <cac:InvoiceLine><cbc:ID>1</cbc:ID><cbc:InvoicedQuantity unitCode="NIU">3</cbc:InvoicedQuantity>
    <cbc:LineExtensionAmount currencyID="PEN">1800.00</cbc:LineExtensionAmount>
    <cac:Item><cbc:Description>Monitor LED 27 pulgadas QHD</cbc:Description></cac:Item>
    <cac:Price><cbc:PriceAmount currencyID="PEN">600.00</cbc:PriceAmount></cac:Price></cac:InvoiceLine>
</Invoice>
"""
    path.write_text(xml, encoding="utf-8")


def make_eml(path: Path, attachment: bytes) -> None:
    msg = EmailMessage()
    msg["From"] = "Andes Cloud Solutions <facturacion@andescloud.pe>"
    msg["To"] = "pagos@portalesdelsur.pe"
    msg["Subject"] = "Factura electrónica F001-00004821"
    msg["Date"] = "Fri, 14 Aug 2026 10:15:00 -0500"
    msg.set_content(
        "Estimado cliente,\n\nAdjuntamos la factura electrónica F001-00004821.\n\nSaludos,\nÁrea de Facturación"
    )
    msg.add_attachment(attachment, maintype="application", subtype="pdf", filename="F001-00004821.pdf")
    path.write_bytes(bytes(msg))


def main() -> None:
    OUT.mkdir(exist_ok=True)

    factura_pdf = html_to_pdf(FACTURA_HTML, fitz.paper_size("a4"))
    (OUT / "factura-electronica.pdf").write_bytes(factura_pdf)

    boleta = scanned_look(html_to_pdf(BOLETA_HTML, (260, 400)), dpi=220, angle=-2.2)
    boleta.save(OUT / "boleta-foto.jpg", quality=80)

    scanned = scanned_look(factura_pdf)
    buffer = BytesIO()
    scanned.save(buffer, format="JPEG", quality=70)
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buffer.getvalue())
    (OUT / "factura-escaneada.pdf").write_bytes(pdf.tobytes(deflate=True))

    make_docx(OUT / "invoice-northwind.docx")
    make_xlsx(OUT / "factura-mar-azul.xlsx")
    make_xml(OUT / "factura-ubl-sunat.xml")
    make_eml(OUT / "email-con-factura.eml", factura_pdf)

    for p in sorted(OUT.iterdir()):
        print(f"{p.name:32} {p.stat().st_size / 1024:8.1f} KB")


if __name__ == "__main__":
    main()
