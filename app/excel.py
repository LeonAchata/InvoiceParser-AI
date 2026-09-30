"""Export an extracted document to a styled Excel workbook."""

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ACCENT = "4F46E5"
LIGHT = "EEF2FF"
THIN = Side(style="thin", color="D4D4D8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_workbook(data: dict[str, Any], filename: str | None = None) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Document"
    ws.sheet_view.showGridLines = False
    currency = data.get("currency") or ""
    money = f'"{currency} "#,##0.00' if currency else "#,##0.00"

    def section(row: int, title: str) -> int:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", start_color=ACCENT)
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 22
        return row + 1

    def pair(row: int, label: str, value: Any, fmt: str | None = None) -> int:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, color="52525B")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=2, value=value if value not in (None, "") else "—")
        if fmt and isinstance(value, (int, float)):
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal="left")
        return row + 1

    ws.merge_cells("A1:E1")
    ws["A1"] = "Extracted document"
    ws["A1"].font = Font(bold=True, size=16, color=ACCENT)
    ws["A2"] = f"Source: {filename or '—'}   ·   Exported {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(italic=True, color="71717A", size=9)

    row = section(4, "DOCUMENT")
    row = pair(row, "Type", data.get("document_type"))
    row = pair(row, "Number", data.get("document_number"))
    row = pair(row, "Issue date", data.get("issue_date"))
    row = pair(row, "Due date", data.get("due_date"))
    row = pair(row, "Currency", currency)
    row = pair(row, "Payment method", data.get("payment_method"))

    for key, title in (("issuer", "ISSUER"), ("customer", "CUSTOMER")):
        party = data.get(key) or {}
        row = section(row + 1, title)
        row = pair(row, "Tax ID", party.get("tax_id"))
        row = pair(row, "Name", party.get("name"))
        row = pair(row, "Address", party.get("address"))
        row = pair(row, "City / district", party.get("city"))

    row = section(row + 1, "LINE ITEMS")
    for col, header in enumerate(("#", "Description", "Quantity", "Unit price", "Amount"), 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True, color="3730A3")
        cell.fill = PatternFill("solid", start_color=LIGHT)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1
    items = data.get("items") or []
    for n, item in enumerate(items, 1):
        values = (n, item.get("description"), item.get("quantity"), item.get("unit_price"), item.get("amount"))
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = BORDER
            if col >= 4:
                cell.number_format = money
            if col == 2:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    if not items:
        ws.cell(row=row, column=2, value="No line items").font = Font(italic=True, color="A1A1AA")
        row += 1

    row += 1
    totals = [("Subtotal", data.get("subtotal")), ("Discount", data.get("discount"))]
    rate = data.get("tax_rate")
    totals.append((f"Tax ({rate:g}%)" if rate else "Tax", data.get("tax")))
    totals.append(("TOTAL", data.get("total")))
    withholding = data.get("withholding") or {}
    if withholding.get("amount"):
        pct = withholding.get("percentage")
        totals.append((f"Withholding ({pct:g}%)" if pct else "Withholding", withholding.get("amount")))
    for label, value in totals:
        if value is None and label == "Discount":
            continue
        ws.cell(row=row, column=4, value=label).font = Font(bold=True)
        cell = ws.cell(row=row, column=5, value=value)
        cell.number_format = money
        if label == "TOTAL":
            for col in (4, 5):
                ws.cell(row=row, column=col).fill = PatternFill("solid", start_color=LIGHT)
                ws.cell(row=row, column=col).font = Font(bold=True, size=12, color="3730A3")
        row += 1

    if data.get("notes"):
        row = section(row + 1, "NOTES")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        ws.cell(row=row, column=1, value=data["notes"]).alignment = Alignment(wrap_text=True)

    for col, width in zip("ABCDE", (18, 48, 12, 16, 18)):
        ws.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
